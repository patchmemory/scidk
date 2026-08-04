"""``TabularFilePlugin`` — the Pipeline's built-in local-file data source.

Task B of the source flow offers two ways to connect: an rclone remote path,
served by whichever plugin owns that remote, and a **file upload**, which has no
plugin behind it. This is that missing half: a
:class:`~scidk.pipeline.plugin_base.DataSourcePlugin` over a CSV/TSV/Excel file
already on local disk, so an uploaded spreadsheet reaches the mapping engine
through exactly the same four calls a SharePoint list does.

It lives in ``scidk/pipeline/`` and not in ``plugins/`` deliberately. A plugin is
an *extension* — something a deployment adds. Reading a local delimited file is
not an extension; it is the floor the Pipeline needs in order to be testable and
to accept an upload at all. Registering it as a plugin would also make it
disableable, and a deployment that switched it off would break file upload with
no indication why.

Scope: what an upload needs and no more — a local path, a header row, lazy
iteration. No rclone, no remotes, no globs. Those belong to a plugin.
"""
from __future__ import annotations

import csv
import logging
import os
from typing import Any, Callable, Dict, Iterator, List, Optional, Tuple

from .plugin_base import AccessResult, DataSourcePlugin, FindResult

logger = logging.getLogger(__name__)

__all__ = ["EXCEL_SUFFIXES", "TabularFilePlugin", "upload_dir"]

#: Where uploaded source files land when nothing else is configured. Under
#: ``~/.scidk`` rather than the cwd, because gunicorn's working directory is not
#: somewhere a deployment should be accumulating user data.
DEFAULT_UPLOAD_DIR = os.path.join(os.path.expanduser("~"), ".scidk", "pipeline_uploads")

#: Suffix → delimiter for the delimited formats. Anything unlisted is read as
#: comma-separated, which is what an unlabelled export almost always is.
_DELIMITERS = {".csv": ",", ".tsv": "\t", ".tab": "\t", ".txt": ","}

#: Suffixes read through openpyxl rather than :mod:`csv`.
EXCEL_SUFFIXES = (".xlsx", ".xlsm")

#: Rows counted before :meth:`TabularFilePlugin.find` gives up on an exact total,
#: so discovery stays inside its 10-second budget on a large file.
MAX_SCAN_ROWS = 100_000


def upload_dir(app: Optional[Any] = None, create: bool = False) -> str:
    """Directory uploaded source files are stored in.

    Precedence: app config, then ``SCIDK_PIPELINE_UPLOAD_DIR``, then
    :data:`DEFAULT_UPLOAD_DIR` — the standard SciDK order. Resolved through
    ``realpath`` so it can be compared against an uploaded file's path without a
    symlink making an outside path look inside.

    Args:
        app: Flask app, when there is one. Scheduled runs pass None and fall back
            to the environment.
        create: Create the directory if absent. Only the upload route needs this;
            a read path should not conjure the directory into existence.
    """
    configured = None
    if app is not None:
        try:
            configured = app.config.get("SCIDK_PIPELINE_UPLOAD_DIR")
        except Exception:  # noqa: BLE001 - a config-less app object
            configured = None
    path = (
        configured
        or os.environ.get("SCIDK_PIPELINE_UPLOAD_DIR")
        or DEFAULT_UPLOAD_DIR
    )
    resolved = os.path.realpath(os.path.expanduser(str(path)))
    if create:
        os.makedirs(resolved, exist_ok=True)
    return resolved


class TabularFilePlugin(DataSourcePlugin):
    """A CSV, TSV or Excel file on local disk.

    Args:
        base_dir: When set, every resolved path must live inside it. The upload
            route passes the uploads directory, so a crafted ``source_path`` in a
            saved source record cannot turn a pipeline run into an arbitrary file
            read. Unset (the default) allows any readable path, which is what
            tests and an admin configuring a local export want.
    """

    name = "tabular_file"
    display_name = "Local file (CSV / TSV / Excel)"
    source_types = ["csv", "tsv", "excel", "file_upload"]

    def __init__(self, base_dir: Optional[str] = None):
        self._base_dir = os.path.realpath(base_dir) if base_dir else None

    # ------------------------------------------------------------ config

    def resolve_source(self, config: Dict[str, Any]) -> str:
        """Return the configured path, or "" when none is set."""
        for key in ("source_path", "file_path", "path", "source"):
            value = str((config or {}).get(key) or "").strip()
            if value:
                return value
        return ""

    def _resolve_checked(self, config: Dict[str, Any]) -> Tuple[str, Optional[str]]:
        """Return ``(realpath, error)``. ``error`` is set when the path is unusable."""
        target = self.resolve_source(config)
        if not target:
            return "", "No file configured. Set 'source_path' to a local file path."
        real = os.path.realpath(os.path.expanduser(target))
        if self._base_dir is not None:
            # commonpath, not startswith: "/uploads-evil" must not pass for "/uploads".
            try:
                inside = os.path.commonpath([real, self._base_dir]) == self._base_dir
            except ValueError:  # different drives on Windows
                inside = False
            if not inside:
                return real, f"{target!r} is outside the permitted directory"
        if not os.path.exists(real):
            return real, f"{target!r} does not exist"
        if not os.path.isfile(real):
            return real, f"{target!r} is not a file"
        return real, None

    @staticmethod
    def _sheet(config: Dict[str, Any]) -> Optional[str]:
        sheet = (config or {}).get("sheet") or (config or {}).get("sheet_name")
        return str(sheet) if sheet not in (None, "") else None

    @staticmethod
    def _positive_int(config: Dict[str, Any], key: str, default: int) -> int:
        try:
            return max(0, int((config or {}).get(key, default)))
        except (TypeError, ValueError):
            return default

    # ----------------------------------------------------- FAIR: Findable

    def find(self, config: Dict[str, Any]) -> FindResult:
        """Read the header and a few rows.

        One lazy pass keeping only the sample, so memory stays flat regardless of
        file size. Counting stops at :data:`MAX_SCAN_ROWS`, after which
        ``row_count`` is None and ``metadata['row_count_truncated']`` is set.
        """
        empty: FindResult = {"ok": False, "columns": [], "row_count": None,
                             "sample": [], "metadata": {}}
        real, error = self._resolve_checked(config)
        if error:
            return {**empty, "error": error}

        sample_rows = self._positive_int(config, "sample_rows", 3)
        max_scan = self._positive_int(config, "max_scan_rows", MAX_SCAN_ROWS) or MAX_SCAN_ROWS
        metadata: Dict[str, Any] = {
            "source_path": real,
            "format": "excel" if real.lower().endswith(EXCEL_SUFFIXES) else "delimited",
            "size_bytes": os.path.getsize(real),
            "mtime": os.path.getmtime(real),
            "transport": "local",
        }
        sheet = self._sheet(config)
        if sheet:
            metadata["sheet"] = sheet

        try:
            columns, rows, close = self._open(real, sheet)
        except Exception as e:  # noqa: BLE001 - an unreadable file is a result
            logger.warning("TabularFilePlugin.find failed for %r: %s", real, e)
            return {**empty, "metadata": metadata, "error": str(e)}

        # No header row means no columns, and a source with no columns cannot be
        # mapped onto anything. Reported as a failure rather than an empty
        # success: a header-only file is a legitimate source that describes its
        # schema, but a file with no header at all describes nothing.
        if not columns:
            return {
                **empty, "metadata": metadata,
                "error": "the file has no header row, so it has no columns to map",
            }

        sample: List[Dict[str, Any]] = []
        count = 0
        truncated = False
        try:
            for row in rows:
                if len(sample) < sample_rows:
                    sample.append(dict(row))
                count += 1
                if count >= max_scan:
                    truncated = True
                    break
        except Exception as e:  # noqa: BLE001 - a malformed file is a result
            return {**empty, "metadata": metadata, "error": str(e)}
        finally:
            close()

        if truncated:
            metadata["row_count_truncated"] = True
        return {
            "ok": True,
            "columns": columns,
            "row_count": None if truncated else count,
            "sample": sample,
            "metadata": metadata,
            "error": None,
        }

    # --------------------------------------------------- FAIR: Accessible

    def access(self, config: Dict[str, Any]) -> AccessResult:
        """Verify the process can actually read the file's bytes.

        ``find()`` succeeding is not the same assurance: a directory can be
        listable while its contents are not readable. This opens the file.
        """
        real, error = self._resolve_checked(config)
        if error:
            return {"ok": False, "auth_method": None, "error": error}
        try:
            with open(real, "rb") as fh:
                fh.read(1)
        except OSError as e:
            return {"ok": False, "auth_method": "local", "error": str(e)}
        return {"ok": True, "auth_method": "local", "error": None}

    # ------------------------------------------------ FAIR: Interoperable

    def fetch(self, config: Dict[str, Any]) -> Iterator[Dict[str, Any]]:
        """Stream the file's rows lazily as raw ``{column: value}`` dicts.

        Raises:
            ValueError: The path is missing, outside the permitted directory, or
                not a readable file. Raised eagerly rather than on first
                iteration, so a misconfiguration surfaces at the call site.
        """
        real, error = self._resolve_checked(config)
        if error:
            raise ValueError(error)
        return self._iter_rows(real, self._sheet(config))

    def _iter_rows(self, real: str, sheet: Optional[str]) -> Iterator[Dict[str, Any]]:
        _columns, rows, close = self._open(real, sheet)
        try:
            yield from rows
        finally:
            close()

    # --------------------------------------------------- FAIR: Reproducible

    def transform_library(self) -> Dict[str, Callable]:
        """No file-specific transforms: a CSV cell is just text.

        Everything a delimited file needs — ``date_parse``, ``boolean_coerce``,
        ``split_delimiter`` and the rest — is source-agnostic and already in
        :mod:`scidk.pipeline.transforms`.
        """
        return {}

    # ---------------------------------------------------------- reading

    def _open(
        self, real: str, sheet: Optional[str]
    ) -> Tuple[List[str], Iterator[Dict[str, Any]], Callable[[], None]]:
        """Return ``(columns, lazy rows, close)``.

        Columns come back separately so a header-only file still describes its
        schema — ``find()`` can report the columns of a file it never reads a row
        of.
        """
        if real.lower().endswith(EXCEL_SUFFIXES):
            return self._open_excel(real, sheet)
        return self._open_delimited(real)

    @staticmethod
    def _open_delimited(
        real: str,
    ) -> Tuple[List[str], Iterator[Dict[str, Any]], Callable[[], None]]:
        delimiter = _DELIMITERS.get(os.path.splitext(real)[1].lower(), ",")
        # utf-8-sig strips the BOM Excel writes; errors='replace' keeps one bad
        # byte from aborting a run that is otherwise fine.
        handle = open(real, "r", encoding="utf-8-sig", errors="replace", newline="")
        try:
            reader = csv.reader(handle, delimiter=delimiter)
            header = next(reader, [])
            columns = _normalize_header(header)

            def rows() -> Iterator[Dict[str, Any]]:
                for values in reader:
                    if not any(str(v).strip() for v in values):
                        continue  # a trailing blank line is not a row
                    yield _zip_row(columns, values)

            return columns, rows(), handle.close
        except Exception:
            handle.close()
            raise

    @staticmethod
    def _open_excel(
        real: str, sheet: Optional[str]
    ) -> Tuple[List[str], Iterator[Dict[str, Any]], Callable[[], None]]:
        from openpyxl import load_workbook

        # read_only + values_only: openpyxl otherwise builds a cell object per
        # cell, which is what turns a 200k-row workbook into gigabytes.
        workbook = load_workbook(real, read_only=True, data_only=True)
        try:
            if sheet:
                if sheet not in workbook.sheetnames:
                    raise ValueError(
                        f"worksheet {sheet!r} not found; this workbook has "
                        f"{workbook.sheetnames}"
                    )
                worksheet = workbook[sheet]
            else:
                worksheet = workbook[workbook.sheetnames[0]]

            stream = worksheet.iter_rows(values_only=True)
            columns = _normalize_header(list(next(stream, ()) or ()))

            def rows() -> Iterator[Dict[str, Any]]:
                for values in stream:
                    values = list(values or ())
                    if not any(v is not None and str(v).strip() for v in values):
                        continue
                    yield _zip_row(columns, values)

            return columns, rows(), workbook.close
        except Exception:
            workbook.close()
            raise


def _normalize_header(header: Any) -> List[str]:
    """Clean a header row into usable column names.

    Internal whitespace is collapsed, matching what the SharePoint plugin does,
    so ``'PI  Archived'`` and ``'PI Archived'`` are the same column and a mapping
    config does not have to guess how many spaces the export used. An unnamed
    column becomes ``column_<n>`` so its values are still addressable.
    """
    columns: List[str] = []
    for index, raw in enumerate(header or ()):
        text = "" if raw is None else " ".join(str(raw).split())
        columns.append(text or f"column_{index + 1}")
    return columns


def _zip_row(columns: List[str], values: List[Any]) -> Dict[str, Any]:
    """Pair a value row with the header.

    A short row leaves the trailing columns as ``None`` (the cell is absent, not
    blank); a long row keeps its extras under ``column_<n>`` rather than dropping
    data the header failed to describe.
    """
    row: Dict[str, Any] = {}
    for index, name in enumerate(columns):
        row[name] = values[index] if index < len(values) else None
    for index in range(len(columns), len(values)):
        row[f"column_{index + 1}"] = values[index]
    return row
