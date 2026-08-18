"""Row readers — turn a source's bytes into ``{column: value}`` dicts.

Format parsing only. These functions take text lines or a file handle and know
nothing about where the bytes came from or what the columns mean: locating and
opening a source is :mod:`plugins.sharepoint_intake.ingest`, and interpreting the
columns is :mod:`scidk.pipeline`.

Kept import-light (stdlib plus a lazy openpyxl) so row parsing is unit-testable
without rclone, Flask, or Neo4j.
"""
from __future__ import annotations

import csv
import os
from typing import Any, Dict, Iterable, Iterator, List, Optional, Tuple

#: Field delimiter by file extension. Anything unlisted is treated as comma.
DELIMITERS = {".csv": ",", ".tsv": "\t", ".tab": "\t"}

#: Extensions read as workbooks rather than delimited text.
EXCEL_EXTS = (".xlsx", ".xlsm")


def source_extension(source: str) -> str:
    """Lowercased file extension of ``source`` ("" when it has none)."""
    return os.path.splitext(str(source or "").split("?", 1)[0])[1].lower()


def is_excel(source: str) -> bool:
    """True when ``source`` names a workbook rather than a delimited export."""
    return source_extension(source) in EXCEL_EXTS


def delimiter_for(source: str) -> str:
    """Field delimiter implied by ``source``'s extension (comma by default)."""
    return DELIMITERS.get(source_extension(source), ",")


def normalize_header(name: Any) -> str:
    """Collapse internal whitespace: ``"PI  Archived"`` -> ``"PI Archived"``.

    SharePoint exports are inconsistent about double spaces; a mapping config
    should not have to be.
    """
    return " ".join(str("" if name is None else name).split())


def clean(value: Any) -> str:
    """Render a raw cell as a stripped string ("" for a missing value)."""
    return "" if value is None else str(value).strip()


def decode_line(chunk: Any) -> str:
    """Decode source bytes, tolerating a BOM and undecodable bytes."""
    if isinstance(chunk, str):
        return chunk
    return bytes(chunk).decode("utf-8-sig", errors="replace")


def dict_rows(columns: List[str], values: Iterable[Any]) -> Iterator[Dict[str, str]]:
    """Zip successive value tuples against ``columns``, skipping blank rows.

    Values past the header width are dropped — the header is authoritative.
    """
    for row in values:
        cleaned = [clean(v) for v in row]
        if not any(cleaned):
            continue
        yield {col: val for col, val in zip(columns, cleaned) if col}


def delimited_rows(lines: Iterable[str],
                   delimiter: str = ",") -> Tuple[List[str], Iterator[Dict[str, str]]]:
    """Return ``(columns, lazy rows)`` for delimited text.

    ``csv.reader`` owns the line pull, so a quoted field spanning newlines parses
    correctly and nothing beyond the current row is held.
    """
    reader = csv.reader(lines, delimiter=delimiter)
    columns = [normalize_header(c) for c in next(reader, [])]
    return columns, dict_rows(columns, reader)


def excel_rows(handle: Any,
               sheet: Optional[str] = None) -> Tuple[List[str], Iterator[Dict[str, str]]]:
    """Return ``(columns, lazy rows)`` for a workbook path or file-like object.

    Read-only mode streams the sheet row by row, but the container itself must be
    fully available: an xlsx is a zip and needs random access.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(handle, read_only=True, data_only=True)
    worksheet = workbook[sheet] if sheet else workbook[workbook.sheetnames[0]]
    values = worksheet.iter_rows(values_only=True)
    columns = [normalize_header(c) for c in (next(values, ()) or ())]

    def _rows() -> Iterator[Dict[str, str]]:
        try:
            yield from dict_rows(columns, values)
        finally:
            workbook.close()

    return columns, _rows()
